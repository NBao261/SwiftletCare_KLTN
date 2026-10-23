import copy
from docx.oxml.ns import qn
GREY='666666'
def _strip_color(r):
    rPr=r._r.rPr
    if rPr is not None:
        for c in rPr.findall(qn('w:color')): rPr.remove(c)
def set_cell(cell, text, bold=None):
    ps=cell.paragraphs
    for p in ps[1:]: p._p.getparent().remove(p._p)
    p=ps[0]
    if p.runs:
        r=p.runs[0]
        for x in p.runs[1:]: x._r.getparent().remove(x._r)
    else:
        r=p.add_run()
    r.text=text
    if r.font.color is not None and r.font.color.type is not None and str(r.font.color.rgb)==GREY: _strip_color(r)
    r.italic=False if r.italic else r.italic
    if bold is not None: r.bold=bold
def fill_table(t, rows, template_row=1, header_rows=1):
    tmpl=copy.deepcopy(t.rows[template_row]._tr)
    for r in list(t.rows)[header_rows:]: r._tr.getparent().remove(r._tr)
    for vals in rows:
        tr=copy.deepcopy(tmpl); t._tbl.append(tr)
        row=t.rows[-1]
        for c,v in zip(row.cells, vals): set_cell(c, v)
def set_kv(t, mapping):
    for r in t.rows:
        k=r.cells[0].text.strip()
        if k in mapping: set_cell(r.cells[1], mapping[k])
def drop_grey(doc):
    for p in list(doc.paragraphs):
        cols={str(r.font.color.rgb) for r in p.runs if r.font.color is not None and r.font.color.type is not None}
        if p.runs and cols=={GREY}: p._p.getparent().remove(p._p)
