"""Write docs/reports/registration-v3.md from the v3 Word file.

Usage: python3 docs/reports/tools/registration_md.py <v3.docx> docs/reports/registration-v3.md
"""
import docx,sys
d=docx.Document(sys.argv[1])
out=["# Capstone Project Register — SwiftletCare (proposed final version v3)","",
"> Text of the registration form submitted for Review 2 (`registration/FA26_SE_Capstone_Project_Register_SwiftletCare_v3.docx`), kept in Markdown so that the self-review prompt and the reports can cite it. Changes from v2 are listed in the Record of Changes (`bm03-record-of-changes.md`) and in section 4, items (8)–(15). Phone numbers are left out of the repository copy; the team adds them to the submitted file.","",
"| Field | Value |","| --- | --- |","| Duration | July 2026 – December 2026 |","| Profession | Software Engineering |","| Supervisor | `[TEAM: supervisor name]` |","| Current version | v2 (after Review 1) |","| Proposed version | v3 (Review 2) |","",
"## 2. Students","","| # | Full name | Student code | E-mail | Role |","| --- | --- | --- | --- | --- |"]
for n,r in enumerate(d.tables[1].rows[1:],1):
    c=[x.text.strip() for x in r.cells]
    out.append(f"| {n} | {c[1]} | {c[2]} | {c[4]} | {c[5]} |")
out.append("")
started=False
for p in d.paragraphs:
    t=p.text.strip()
    if t.startswith('3. Register content'): started=True
    if not started or not t: continue
    st=p.style.name
    if st=='Heading 2': out+=['## '+t.rstrip(':'),'']
    elif t in ('A. Context:','B. Proposed Solutions:','C. Functional requirement:','D. Main flow:','E. Non-functional requirement:') or t.startswith('3.2.'):
        out+=['### '+t.rstrip(':'),'']
    elif t[:2] in ('1.','2.','3.','4.') and t.endswith(':'): out+=['#### '+t.rstrip(':'),'']
    elif st=='List Paragraph': out.append('- '+t)
    else:
        if out and out[-1].startswith('- '): out.append('')
        runs=p.runs
        if len(runs)>=2 and runs[0].text.strip().endswith(':') and runs[1].text.strip():
            out+=[f"**{runs[0].text.strip()}** {''.join(r.text for r in runs[1:]).strip()}",'']
        elif t.endswith(':') and len(t)<80: out+=['**'+t+'**','']
        else: out+=[t,'']
s='\n'.join(out).replace(':**\n\n- ',':**\n\n- ')
open(sys.argv[2],'w').write(s.rstrip()+'\n')
