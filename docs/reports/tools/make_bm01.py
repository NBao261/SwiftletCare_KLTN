"""Fill form BM01 from docs/reports/design-brief.md.

Usage (repository root): python3 docs/reports/tools/make_bm01.py
"""
import os, sys, re
HERE=os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
REPORTS=os.path.dirname(HERE)
import docx
from docxfill import *
from md_tables import tables
md=os.path.join(REPORTS,'design-brief.md')
T=tables(md)
d=docx.Document(os.path.join(REPORTS,'templates','SWP490_R2_BM01_Design_Brief.docx'))
assert len(T)==12, len(T)
set_kv(d.tables[0], {r[0]:r[1] for r in T[0][1][1:]})
for i in range(1,12):
    fill_table(d.tables[i], T[i][1][1:])
# section 2 text
txt=open(md,encoding='utf-8').read()
glance=re.search(r'## 2\. System at a Glance\n\n(.+?)\n\n', txt, re.S).group(1).replace('`','')
paras=d.paragraphs
i=[k for k,p in enumerate(paras) if p.text.startswith('2. System at a Glance')][0]
target=[p for p in paras[i+1:i+4] if not p.text.strip()][0]
target.add_run(glance)
drop_grey(d)
import copy as _c
def note_after(tbl, text):
    np_=_c.deepcopy(target._p); tbl._tbl.addnext(np_)
    para=docx.text.paragraph.Paragraph(np_, target._parent)
    for r in para.runs[1:]: r._r.getparent().remove(r._r)
    para.runs[0].text=text
note_after(d.tables[5], 'Each gap with its requirements: Project Tracking, Section 6 (G-01…G-19). High findings of the self-review (BM02) marked Kept or Planned are added here after the prompt is run.')
note_after(d.tables[6], 'FO = Farm Owner; DE = Domain Expert (Farm Owner if none, BR-22); auto = done by a system component, not an actor; TC-Fnn = TC-Fnn-01..03 in Report 5. Status is copied from Project Tracking.')
note_after(d.tables[9], 'Actors are people and external services only; the controller, the AI node and the backend are parts of SwiftletCare (SRS 2.3). WF-03 and WF-09 have one actor and are valid because each passes through three mandatory features.')
# A3 note, taken from the Markdown
a3=re.search(r'^Automated in CI:.*$', txt, re.M).group(0)
for p in d.paragraphs:
    if p.text.startswith('Automated in CI'):
        p.runs[0].text=a3
        for r in p.runs[1:]: r.text=''
d.save(os.path.join(REPORTS,'design-brief','SWP490_R2_BM01_Design_Brief_SwiftletCare.docx'))
