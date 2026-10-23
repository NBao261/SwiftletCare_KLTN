"""Fill form BM03 from docs/reports/bm03-record-of-changes.md.

Usage (repository root): python3 docs/reports/tools/make_bm03.py
"""
import os, sys
HERE=os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
REPORTS=os.path.dirname(HERE)
import docx
from docxfill import *
from md_tables import tables
d=docx.Document(os.path.join(REPORTS,'templates','SWP490_R2_BM03_Record_of_Changes.docx'))
t=tables(os.path.join(REPORTS,'bm03-record-of-changes.md'))
info={r[0]:r[1] for r in t[0][1][1:]}
rows=[r[1:] for r in t[1][1][1:]]
set_kv(d.tables[0], info)
fill_table(d.tables[1], [[str(i+1)]+r for i,r in enumerate(rows)])
fill_table(d.tables[2], [[str(i+1),'',''] for i in range(len(rows))])
drop_grey(d)
d.save(os.path.join(REPORTS,'registration','SWP490_R2_BM03_Record_of_Changes_SwiftletCare.docx'))
